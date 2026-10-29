from pathlib import Path
import openpyxl,json,re,zipfile,xml.etree.ElementTree as ET,math
import formulas
root=Path(__file__).resolve().parents[1];file=root/'public/downloads/FREEDOM_MIRROR_X_2026_PRELIMINARY.xlsx'
calc=formulas.ExcelModel().loads(str(file)).finish().calculate();cached=openpyxl.load_workbook(file,data_only=True);formula_wb=openpyxl.load_workbook(file,data_only=False)
lookup={str(k).upper():v for k,v in calc.items()};checks=[];errors=[]
for sheet in formula_wb:
 for row in sheet:
  for cell in row:
   expected=cached[sheet.title][cell.coordinate].value
   if cell.data_type!='f' or not isinstance(expected,(int,float)):continue
   suffix=f"]{sheet.title.upper()}'!{cell.coordinate.upper()}";matches=[v for k,v in lookup.items() if k.endswith(suffix)]
   if not matches:raise AssertionError(f'No result {sheet.title}!{cell.coordinate}')
   actual=float(matches[0].value[0][0]);assert math.isclose(actual,expected,rel_tol=1e-10,abs_tol=1e-7),(sheet.title,cell.coordinate,actual,expected)
   checks.append({'cell':sheet.title+'!'+cell.coordinate,'independent':actual,'cached':expected})
for k,v in calc.items():
 if any(e in str(v) for e in ['#VALUE!','#REF!','#DIV/0!','#NAME?','#NUM!']):errors.append(str(k)+' '+str(v))
assert not errors,errors
snap=json.loads((root/'data/snapshot.json').read_text());ns={'a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
with zipfile.ZipFile(root/'public/downloads/FREEDOM_MIRROR_X_2026_PRELIMINARY.pptx') as z:
 slide_paths=[n for n in z.namelist() if re.fullmatch(r'ppt/slides/slide\d+\.xml',n)];assert len(slide_paths)==14
 first='\n'.join(e.text or '' for e in ET.fromstring(z.read('ppt/slides/slide1.xml')).findall('.//a:t',ns))
 assert snap['id'] in first and 'PRELIMINARY' in first
 for i,c in enumerate(snap['companies']):
  assert c['ticker'] in first
  assert f"${c['target12m']:,.2f}" in first,(c['ticker'],first)
  assert math.isclose(cached['PORTFOLIO'].cell(i+3,4).value,c['target12m'],rel_tol=1e-12)
  assert cached['PORTFOLIO'].cell(i+3,6).value==c['allocationBps']/10000
assert cached['PORTFOLIO']['G6'].value==100000
report={'snapshotId':snap['id'],'independentEngine':'formulas Python package','numericFormulaCellsMatched':len(checks),'formulaErrors':errors,'presentationSlides':14,'pptxXlsxSnapshotReconciliation':'PASS','status':'PRELIMINARY','checks':checks}
(root/'docs/EXPORT_VERIFICATION.json').write_text(json.dumps(report,ensure_ascii=False,indent=2));print(json.dumps({k:v for k,v in report.items() if k!='checks'},ensure_ascii=False,indent=2))
